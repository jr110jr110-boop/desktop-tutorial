# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
D="/tmp/claude-0/-home-claude/b3dbf4fa-9363-5b97-953d-a1c563004b57/scratchpad/kitfin/PMIS_Finance_Test/3-Data/"
FA=str.maketrans("0123456789","۰۱۲۳۴۵۶۷۸۹")
def mk(name, header, rows, widths, sheet="Sheet1"):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title=sheet; ws.sheet_view.rightToLeft=True
    ws.append(header)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font=Font(name="B Zar",bold=True,size=12); c.fill=PatternFill("solid",fgColor="DDDDDD"); c.alignment=Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name="B Zar",size=12)
            if isinstance(c.value,(int,float)): c.number_format='#,##0'
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    ws.freeze_panes="A2"
    wb.save(D+name)

mk("details-sample.xlsx",["کد تفصیلی","عنوان","نوع"],
   [["S-001","تأمین‌کننده آزمایشی ۱","تأمین‌کننده"],["C-001","مشتری آزمایشی ۱","مشتری"]],[16,34,16])
mk("opening-balances-sample.xlsx",["کد حساب","کد تفصیلی","بدهکار","بستانکار"],
   [["10101",None,20000000,0],["10102",None,500000000,0],["10301","C-001",300000000,0],["10501",None,200000000,0],
    ["30101","S-001",0,150000000],["50101",None,0,600000000],["50301",None,0,270000000]],[14,14,18,18])
mk("opening-balances-FAULTY.xlsx",["کد حساب","کد تفصیلی","بدهکار","بستانکار"],
   [["10102",None,500000000,0],["50101",None,0,400000000],["99999",None,0,90000000]],[14,14,18,18])
mk("mass-vouchers-sample.xlsx",["گروه","تاریخ","شرح سند","کد حساب","کد تفصیلی اول","کد تفصیلی دوم","بدهکار","بستانکار","شرح ردیف"],
   [[1,"۱۴۰۵/۰۸/۰۴","هزینه آزمایشی ۱","80108",None,None,1500000,0,"هزینه"],
    [1,"۱۴۰۵/۰۸/۰۴","هزینه آزمایشی ۱","10101",None,None,0,1500000,"پرداخت از صندوق"],
    [2,"۱۴۰۵/۰۸/۰۴","هزینه آزمایشی ۲","80107",None,None,2000000,0,"هزینه"],
    [2,"۱۴۰۵/۰۸/۰۴","هزینه آزمایشی ۲","10101",None,None,0,2000000,"پرداخت از صندوق"],
    [3,"۱۴۰۵/۰۸/۰۴","سند نامتوازن (باید رد شود)","80108",None,None,700000,0,"هزینه"],
    [3,"۱۴۰۵/۰۸/۰۴","سند نامتوازن (باید رد شود)","10101",None,None,0,600000,"پرداخت از صندوق"]],[8,14,28,12,14,14,14,14,22])
# bank statement: opening 0, closing 217,950,000
bs=[("۱۴۰۵/۰۸/۰۱","واریز اولیه مشتری آزمایشی ۳",0,100000000,"RCP-T0"),
    ("۱۴۰۵/۰۸/۰۹","شارژ تنخواه",3000000,0,"PAY-T3"),
    ("۱۴۰۵/۰۸/۱۰","پرداخت به تأمین‌کننده آزمایشی ۲",9900000,0,"PAY-T1"),
    ("۱۴۰۵/۰۸/۱۱","پرداخت به تأمین‌کننده آزمایشی ۳",18400000,0,"PAY-T2"),
    ("۱۴۰۵/۰۸/۱۲","واریز مشتری آزمایشی ۲ (پیش‌دریافت)",0,55000000,"RCP-T1"),
    ("۱۴۰۵/۰۸/۱۴","واریز مشتری آزمایشی ۲ (صورت‌وضعیت)",0,72000000,"RCP-T2"),
    ("۱۴۰۵/۰۸/۱۵","واریز مشتری آزمایشی ۲ (آزادسازی)",0,10000000,"RCP-T3"),
    ("۱۴۰۵/۰۸/۱۸","وصول چک A",0,20000000,"CHQ-A"),
    ("۱۴۰۵/۰۸/۲۰","پاس شدن چک شماره 1001",8000000,0,"1001"),
    ("۱۴۰۵/۰۸/۲۰","کارمزد بانکی",150000,0,"FEE-1"),
    ("۱۴۰۵/۰۸/۲۰","سود سپرده",0,400000,"INT-1")]
bal=0; rows=[]
for d,desc,w,dep,ref in bs:
    bal+=dep-w; rows.append([d,desc,w,dep,ref,bal])
assert bal==217950000, bal
mk("bank-statement-sample.xlsx",["تاریخ","شرح","برداشت","واریز","مرجع","مانده"],rows,[14,38,14,14,12,16])
mk("supplier-statement-sample.xlsx",["تاریخ","شماره سند","شرح","بستانکار","بدهکار"],
   [["۱۴۰۵/۰۸/۰۴","INV-T1","فاکتور فروش پیچ",11000000,0],
    ["۱۴۰۵/۰۸/۰۶","CN-1","برگشت از خرید ۱۰ عدد",0,1100000],
    ["۱۴۰۵/۰۸/۱۰","PAY-T1","دریافت وجه",0,9900000],
    ["۱۴۰۵/۰۸/۱۵","INV-X9","فاکتور ثبت‌نشده در PMIS",500000,0]],[14,14,34,16,16])
print("ok")
