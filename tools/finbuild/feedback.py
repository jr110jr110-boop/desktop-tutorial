# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from build import T, SEC, QUESTIONS, LINKED, plain
OUT="/tmp/claude-0/-home-claude/b3dbf4fa-9363-5b97-953d-a1c563004b57/scratchpad/kitfin/PMIS_Finance_Test/4-Feedback/Test-Feedback-Finance.xlsx"
F="B Zar"
thin=Side(style="thin",color="999999"); BR=Border(left=thin,right=thin,top=thin,bottom=thin)
HF=PatternFill("solid",fgColor="DDDDDD")
wb=openpyxl.Workbook()

def head(ws,cols,widths):
    ws.append(cols)
    for c in ws[1]:
        c.font=Font(name=F,bold=True,size=12); c.fill=HF; c.border=BR
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for i,w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
    ws.freeze_panes="A2"; ws.sheet_view.rightToLeft=True
    ws.row_dimensions[1].height=34

def body(ws,first_row,ncols,wrap_cols=None):
    for row in ws.iter_rows(min_row=first_row,max_row=ws.max_row,max_col=ncols):
        for c in row:
            c.font=Font(name=F,size=11); c.border=BR
            c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="right")

# ---- sheet 1
ws=wb.active; ws.title="نتایج تست"
cols=["شناسه","بخش","سناریو","نقش","نتیجه مورد انتظار","نتیجه (OK/NOK)","مشاهده واقعی / توضیح","شماره تصویر","شدت","سمت آزمون‌کننده","تاریخ (شمسی)"]
head(ws,cols,[8,22,28,22,60,14,44,12,12,20,14])
for x in T:
    ws.append([x["id"],SEC[x["sec"]],plain(x["title"]),plain(x["role"]),plain(x["exp"]),None,None,None,None,None,None])
n=ws.max_row
body(ws,2,11)
for r in range(2,n+1):
    ws.cell(r,1).alignment=Alignment(horizontal="center",vertical="top")
    ws.cell(r,1).font=Font(name=F,size=11,bold=True)
    for c in (6,8,9,11): ws.cell(r,c).alignment=Alignment(horizontal="center",vertical="top",wrap_text=True)
    for c in (6,7,8,9,10,11): ws.cell(r,c).fill=PatternFill("solid",fgColor="FFFFF2")
dv=DataValidation(type="list",formula1='"OK,NOK"',allow_blank=True,showErrorMessage=True,errorTitle="مقدار نامعتبر",error="فقط OK یا NOK")
dv.promptTitle="نتیجه"; dv.prompt="OK = مطابق نتیجه مورد انتظار؛ NOK = نادرست یا انجام نشد (علت را در ستون توضیح بنویسید)"; dv.showInputMessage=True
dv2=DataValidation(type="list",formula1='"بحرانی,مهم,جزئی"',allow_blank=True,showErrorMessage=True,errorTitle="مقدار نامعتبر",error="بحرانی، مهم یا جزئی")
ws.add_data_validation(dv); ws.add_data_validation(dv2)
dv.add("F2:F%d"%n); dv2.add("I2:I%d"%n)
green=PatternFill("solid",bgColor="C6EFCE",fgColor="C6EFCE"); red=PatternFill("solid",bgColor="FFC7CE",fgColor="FFC7CE")
ws.conditional_formatting.add("F2:F%d"%n,CellIsRule(operator="equal",formula=['"OK"'],fill=green))
ws.conditional_formatting.add("F2:F%d"%n,CellIsRule(operator="equal",formula=['"NOK"'],fill=red))
ws.auto_filter.ref="A1:K%d"%n
ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
ws.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
ws.print_title_rows="1:1"

# ---- sheet 2
w2=wb.create_sheet("سؤالات حسابدار")
head(w2,["شماره","سؤال","آزمون مرتبط","پاسخ","تصمیم‌گیرنده (عنوان شغلی)","تاریخ (شمسی)"],[8,70,14,60,24,14])
for i,q in enumerate(QUESTIONS,1):
    w2.append([int(plain(str(i)).translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹","0123456789"))) if False else i,plain(q),plain(LINKED[i]).replace("، ","\n"),None,None,None])
body(w2,2,6)
for r in range(2,w2.max_row+1):
    w2.cell(r,1).alignment=Alignment(horizontal="center",vertical="top")
    w2.cell(r,1).number_format="0"
    w2.row_dimensions[r].height=95
    for c in (4,5,6): w2.cell(r,c).fill=PatternFill("solid",fgColor="FFFFF2")
w2.page_setup.orientation="landscape"; w2.page_setup.fitToWidth=1; w2.page_setup.fitToHeight=0
w2.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)

# ---- sheet 3
w3=wb.create_sheet("راهنما")
w3.sheet_view.rightToLeft=True; w3.column_dimensions["A"].width=26; w3.column_dimensions["B"].width=100
rows=[
 ("هدف","ثبت نتیجه آزمون بسته حسابداری مالی PMIS. هر ردیف برگه «نتایج تست» یک آزمون از راهنمای PDF (0-Test-Guide-Finance-FA.pdf) است و کد آن (F-01 تا F-71) همان کد راهنماست."),
 ("پایگاه آزمون","همه کارها فقط روی پایگاه آزمایشی انجام می‌شود، با داده و نام آزمایشی. دفتر واقعی، حقوق واقعی و رمز یا نام کاربری در این فرم نوشته نمی‌شود."),
 ("ستون نتیجه","فقط یکی از OK یا NOK را از فهرست انتخاب کنید. OK یعنی هر آنچه در «نتیجه مورد انتظار» آمده دیده شد (از جمله عددها). NOK یعنی نتیجه فرق داشت، خطا آمد یا آزمون انجام نشد."),
 ("مشاهده واقعی / توضیح","برای هر NOK بنویسید چه دیدید (متن پیام خطا، عدد دیده‌شده در برابر عدد مورد انتظار، منوی محل). برای OK هم اگر نکته‌ای هست بنویسید. اگر آزمون انجام نشد بنویسید «انجام نشد» و علت."),
 ("شماره تصویر","نام یا شماره فایل تصویر (مثلاً F-26-1) که در پوشه 4-Feedback کنار همین فایل گذاشته‌اید."),
 ("شدت","بحرانی: مبلغ یا دفاتر غلط، یا قانون «تهیه‌کننده ≠ تاییدکننده» رعایت نمی‌شود، یا اطلاعات حقوق به کسی غیرمجاز نشان داده می‌شود. مهم: کار انجام نمی‌شود یا خطا می‌دهد. جزئی: متن، ظاهر، جابه‌جایی ستون."),
 ("سمت آزمون‌کننده","عنوان شغلی، نه نام شخص: مثلاً حسابدار ۱، حسابدار ۲، مدیر مالی، خزانه‌داری، حقوق و دستمزد، حسابرس، مدیرعامل."),
 ("تاریخ","تاریخ شمسی به‌صورت ۱۴۰۵/۰۸/۱۰. تاریخ میلادی نوشته نشود."),
 ("تفکیک وظایف","در آزمون‌هایی که دو نفر لازم است، با دو کاربر جدا کار کنید. اگر یک نفر توانست کار تهیه و تایید را هر دو بکند، NOK با شدت بحرانی بزنید."),
 ("برگه سؤالات حسابدار","پاسخ هر یک از ۱۶ سؤال را در ستون «پاسخ» بنویسید، عنوان شغلی تصمیم‌گیرنده و تاریخ را هم پر کنید. این پاسخ‌ها مبنای تنظیم نرخ‌ها و قواعد نسخه بعدی است."),
 ("پس از پایان","فایل را با نام Test-Feedback-Finance-(عنوان شغلی).xlsx ذخیره و برای مدیر فنی بفرستید. تا امضای حسابدار، دفاتر واقعی در PMIS وارد نمی‌شود."),
]
w3.append(["موضوع","توضیح"])
for c in w3[1]: c.font=Font(name=F,bold=True,size=12); c.fill=HF; c.border=BR; c.alignment=Alignment(horizontal="center")
for a,b in rows: w3.append([a,b])
for row in w3.iter_rows(min_row=2):
    for c in row:
        c.font=Font(name=F,size=11,bold=(c.column==1)); c.border=BR; c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="right")
wb.save(OUT); print("saved",n-1,"tests")
