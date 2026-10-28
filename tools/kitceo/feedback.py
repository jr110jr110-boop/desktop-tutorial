# -*- coding: utf-8 -*-
"""Writes the result form of the CEO executive dashboard (release D1) test kit.
Run after build.py: python3 feedback.py
Sheets: «نتایج تست» (one row per scenario), «مفروضات ★», «پرسش‌های شرکت», «نظرهای کلی», «راهنما». All sheets right-to-left, font B Zar."""
import sys, os, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from build import T, SEC, QUESTIONS, LINKED, STAR, STAR_LINKED, plain

OUT = "/home/claude/kitceo/out/PMIS_CEO_Dashboard_Test/4-Feedback/CEO-Dashboard-Test-Feedback.xlsx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)
F = "B Zar"
thin = Side(style="thin", color="999999")
BR = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill("solid", fgColor="DDDDDD")
FILL_IN = PatternFill("solid", fgColor="FFFFF2")
wb = openpyxl.Workbook()


def head(ws, cols, widths):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(name=F, bold=True, size=12)
        c.fill = HF
        c.border = BR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = True
    ws.row_dimensions[1].height = 34


def body(ws, first_row, ncols):
    for row in ws.iter_rows(min_row=first_row, max_row=ws.max_row, max_col=ncols):
        for c in row:
            c.font = Font(name=F, size=11)
            c.border = BR
            c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="right")


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"


def height(pairs, minimum=60, line=15.5, pad=8):
    """Row height that holds the longest wrapped text (pairs of text and column width)."""
    lines = 1
    for text, width in pairs:
        if not text:
            continue
        n = 0
        for part in str(text).split("\n"):
            n += max(1, math.ceil(len(part) / (width * 1.05)))
        lines = max(lines, n)
    return max(minimum, lines * line + pad)


# ---------------------------------------------------------------- sheet 1: results, one row per scenario
ws = wb.active
ws.title = "نتایج تست"
head(ws, ["شناسه", "بخش", "عنوان آزمون", "نقش", "نتیجه (OK / مشکل / انجام نشد)", "توضیح (چه دیدید)", "نام فایل تصویر"],
     [9, 30, 44, 26, 22, 60, 24])
for x in T:
    ws.append([x["id"], SEC[x["sec"]], plain(x["title"]), plain(x["role"]), None, None, None])
n = ws.max_row
body(ws, 2, 7)
for r in range(2, n + 1):
    ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(r, 1).font = Font(name=F, size=11, bold=True)
    ws.cell(r, 5).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for c in (5, 6, 7):
        ws.cell(r, c).fill = FILL_IN
dv = DataValidation(type="list", formula1='"OK,مشکل,انجام نشد"', allow_blank=True, showErrorMessage=True,
                    errorTitle="مقدار نامعتبر", error="فقط OK، مشکل یا انجام نشد")
dv.promptTitle = "نتیجه"
dv.prompt = "OK = مطابق نتیجه مورد انتظار؛ مشکل = نادرست (توضیح بنویسید)؛ انجام نشد = آزمون اجرا نشد (علت را بنویسید)"
dv.showInputMessage = True
ws.add_data_validation(dv)
dv.add("E2:E%d" % n)
green = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
red = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")
grey = PatternFill("solid", bgColor="E0E0E0", fgColor="E0E0E0")
ws.conditional_formatting.add("E2:E%d" % n, CellIsRule(operator="equal", formula=['"OK"'], fill=green))
ws.conditional_formatting.add("E2:E%d" % n, CellIsRule(operator="equal", formula=['"مشکل"'], fill=red))
ws.conditional_formatting.add("E2:E%d" % n, CellIsRule(operator="equal", formula=['"انجام نشد"'], fill=grey))
ws.auto_filter.ref = "A1:G%d" % n
landscape(ws)

# ---------------------------------------------------------------- sheet 2: the ★ assumptions
w2 = wb.create_sheet("مفروضات ★")
W2 = [7, 78, 18, 20, 52, 24, 14]
head(w2, ["ردیف", "مفروضه (قاعده‌ای که شرکت باید تایید کند)", "آزمون مرتبط", "تصمیم (تایید / مقدار دیگر)",
          "مقدار دیگر یا توضیح", "تصمیم‌گیرنده (عنوان شغلی)", "تاریخ (شمسی)"], W2)
for i, s in enumerate(STAR, 1):
    w2.append([i, plain(s), plain(STAR_LINKED[i]).replace("، ", "\n"), None, None, None, None])
body(w2, 2, 7)
for r in range(2, w2.max_row + 1):
    w2.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
    w2.cell(r, 4).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    w2.row_dimensions[r].height = height([(w2.cell(r, 2).value, W2[1]), (w2.cell(r, 3).value, W2[2])])
    for c in (4, 5, 6, 7):
        w2.cell(r, c).fill = FILL_IN
dv2 = DataValidation(type="list", formula1='"تایید,مقدار دیگر"', allow_blank=True, showErrorMessage=True,
                     errorTitle="مقدار نامعتبر", error="فقط «تایید» یا «مقدار دیگر»")
dv2.promptTitle = "تصمیم"
dv2.prompt = "تایید = همین قاعده درست است؛ مقدار دیگر = قاعده یا عدد دیگری می‌خواهیم (در ستون بعد بنویسید)"
dv2.showInputMessage = True
w2.add_data_validation(dv2)
dv2.add("D2:D%d" % w2.max_row)
landscape(w2)

# ---------------------------------------------------------------- sheet 3: the company questions
w3 = wb.create_sheet("پرسش‌های شرکت")
W3 = [8, 74, 18, 62, 26, 14]
head(w3, ["شماره", "پرسش", "آزمون مرتبط", "پاسخ و تصمیم", "تصمیم‌گیرنده (عنوان شغلی)", "تاریخ (شمسی)"], W3)
for i, q in enumerate(QUESTIONS, 1):
    w3.append([i, plain(q), plain(LINKED[i]).replace("، ", "\n"), None, None, None])
body(w3, 2, 6)
for r in range(2, w3.max_row + 1):
    w3.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
    w3.row_dimensions[r].height = height([(w3.cell(r, 2).value, W3[1]), (w3.cell(r, 3).value, W3[2])], minimum=95)
    for c in (4, 5, 6):
        w3.cell(r, c).fill = FILL_IN
landscape(w3)

# ---------------------------------------------------------------- sheet 4: general remarks
w4 = wb.create_sheet("نظرهای کلی")
head(w4, ["موضوع (بخش یا صفحه)", "نظر یا پیشنهاد", "اهمیت (زیاد / متوسط / کم)", "نام فایل تصویر"], [30, 90, 24, 24])
for _ in range(30):
    w4.append([None, None, None, None])
body(w4, 2, 4)
dv4 = DataValidation(type="list", formula1='"زیاد,متوسط,کم"', allow_blank=True)
w4.add_data_validation(dv4)
dv4.add("C2:C31")
for r in range(2, w4.max_row + 1):
    w4.row_dimensions[r].height = 40
    for c in (1, 2, 3, 4):
        w4.cell(r, c).fill = FILL_IN
landscape(w4)

# ---------------------------------------------------------------- sheet 5: how to fill the form
w5 = wb.create_sheet("راهنما")
w5.sheet_view.rightToLeft = True
w5.column_dimensions["A"].width = 26
w5.column_dimensions["B"].width = 110
rows = [
    ("هدف", plain("ثبت نتیجه‌ی آزمون بسته‌ی «داشبورد مدیرعامل» PMIS (انتشار D1). هر ردیف برگه‌ی «نتایج تست» یک آزمون از راهنمای PDF "
                  "(⟦0-Test-Guide-CEO-Dashboard-FA.pdf⟧) است و کد آن (⟦T-01⟧ تا ⟦T-%02d⟧) همان کد راهنماست؛ آزمون‌ها در بخش‌های "
                  "⟦S0⟧ تا ⟦S%d⟧ هستند." % (len(T), len(SEC) - 1))),
    ("پایگاه آزمون", "همه‌ی کارها فقط روی پایگاه و سرور آزمایشی انجام می‌شود، با داده و نام آزمایشی. رمز یا نام کاربری در این فرم نوشته نمی‌شود. "
                     "داشبورد همان داده‌ای را نشان می‌دهد که در پایگاه آزمایشی شما هست؛ عددهای عکس‌های راهنما ساختگی‌اند و با عددهای شما فرق می‌کنند."),
    ("ستون نتیجه", "فقط یکی از OK، مشکل یا انجام نشد را از فهرست انتخاب کنید. OK یعنی هر آنچه در «نتیجه‌ی مورد انتظار» راهنما آمده دیده شد "
                   "(از جمله پیام‌ها، عددها و رنگ‌ها). مشکل یعنی نتیجه فرق داشت یا خطا آمد. انجام نشد یعنی آزمون اجرا نشد (علت را بنویسید)."),
    ("توضیح", "برای هر «مشکل» بنویسید چه دیدید: متن پیام خطا، عدد دیده‌شده در برابر عدد مورد انتظار (و صفحه‌ای که با آن سنجیدید)، رنگ یا نشان کاشی، "
              "و منوی محل. برای «انجام نشد» علت را بنویسید (مثلاً «پایگاه خالی است»، «ماشین میکروفون ندارد»)."),
    ("مشکل بحرانی", "اگر با مشکل زیر روبه‌رو شدید در توضیح اول بنویسید «بحرانی»: عدد داشبورد با صفحه‌ی ماژول خودش برابر نیست؛ نقشی که باید رد شود داده یا "
                    "جدول می‌بیند؛ مبلغ حقوقی که برای معاونت PMO باید «محرمانه (حقوق)» باشد دیده می‌شود؛ دستیار یا صدا کاری می‌کند که تایید انسانی "
                    "می‌خواهد (تایید، نهایی‌کردن، پرداخت، حذف)؛ عددی در پاسخ دستیار که در هیچ نتیجه‌ی سامانه نیست با نشان «راستی‌آزمایی‌شده» می‌آید؛ یا داشبورد باز نمی‌شود."),
    ("نام فایل تصویر", "نام فایل تصویر (مثلاً ⟦T-26-1⟧) که در پوشه‌ی ⟦4-Feedback⟧ کنار همین فایل گذاشته‌اید. در تصویر رمز یا نام کاربری دیده نشود."),
    ("آزمون‌های نشان‌دار", "«(شرطی)» یعنی آزمون فقط وقتی ممکن است که داده‌ی لازم در پایگاه باشد؛ اگر نیست «انجام نشد» بزنید و علت را بنویسید. "
                           "«(ویژه‌ی مدیر فنی)» یعنی دستور SQL یا PowerShell یا رابط برنامه‌نویسی لازم است و مدیر فنی آن را انجام می‌دهد (پیوست ۶ راهنما)."),
    ("مدل آزمایشی", "آزمون‌های «تحلیل»، «دستیار» و «صدا» با مدل آزمایشی (پوشه‌ی ⟦5-Mock-AI⟧) انجام می‌شود که فقط جمله‌های ثابت می‌دهد، گفتار را نمی‌فهمد و صدایش "
                    "سکوت است. رفتارهایی که فقط مدل واقعی نشان می‌دهد در راهنما «با یک مدل واقعی» نوشته شده؛ آن‌ها را «انجام نشد» بزنید و در توضیح بنویسید «مدل واقعی نبود»."),
    ("مفروضات ★", "برگه‌ی «مفروضات ★» قاعده‌هایی است که از ظاهر برنامه و کد درآمده و شرکت باید تایید کند (آستانه‌ی رنگ‌ها، قاعده‌ی هر شاخص، پنهان‌سازی حقوق، چاپ). "
                  "برای هر ردیف «تایید» یا «مقدار دیگر» را انتخاب کنید و اگر مقدار دیگری می‌خواهید در ستون بعد بنویسید. فرق‌داشتنِ ★ با خواسته‌ی شرکت خطا نیست؛ تصمیم شرکت است."),
    ("پرسش‌های شرکت", "پاسخ هر پرسش، عنوان شغلی تصمیم‌گیرنده و تاریخ را بنویسید. اگر پاسخ را هنوز نمی‌دانید «در انتظار تصمیم» و نام واحد تصمیم‌گیرنده را بنویسید."),
    ("نقش‌ها", "عنوان شغلی، نه نام شخص: مدیر سیستم (Admin)، مدیرعامل، معاونت PMO، مدیر مالی، انباردار آزمایشی (و مدیر فنی برای آزمون‌های نشان‌دار)."),
    ("تاریخ", plain("تاریخ شمسی به‌صورت ⟦1405/08/06⟧ نوشته شود. تاریخ میلادی نوشته نشود.")),
    ("پس از پایان", plain("فایل را با نام ⟦CEO-Dashboard-Test-Feedback-⟧ و پس از آن عنوان شغلی خودتان با حروف لاتین ذخیره کنید "
                          "(مثلاً ⟦CEO-Dashboard-Test-Feedback-PMO.xlsx⟧) و برای مدیر فنی بفرستید.")),
]
w5.append(["موضوع", "توضیح"])
for c in w5[1]:
    c.font = Font(name=F, bold=True, size=12)
    c.fill = HF
    c.border = BR
    c.alignment = Alignment(horizontal="center")
for a, b in rows:
    w5.append([a, b])
for row in w5.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=11, bold=(c.column == 1))
        c.border = BR
        c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="right")
for r in range(2, w5.max_row + 1):
    w5.row_dimensions[r].height = height([(w5.cell(r, 2).value, 110)], minimum=34)
w5.page_setup.orientation = "landscape"
w5.page_setup.paperSize = 9
w5.page_setup.fitToWidth = 1
w5.page_setup.fitToHeight = 0
w5.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

wb.save(OUT)
print("saved", OUT, "|", len(T), "tests,", len(STAR), "star rows,", len(QUESTIONS), "questions")
