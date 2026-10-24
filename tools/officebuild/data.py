# -*- coding: utf-8 -*-
import os, csv, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from example import DAYS
D = "/home/claude/kitoffice/out/PMIS_Office_Test/3-Data/"
os.makedirs(D, exist_ok=True)

def mk(name, header, rows, widths, sheet="Sheet1"):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = sheet; ws.sheet_view.rightToLeft = True
    ws.append(header)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font = Font(name="B Zar", bold=True, size=12); c.fill = PatternFill("solid", fgColor="DDDDDD"); c.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="B Zar", size=12)
            c.number_format = "@" if isinstance(c.value, str) else "General"
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(D + name)

# clock import: rows in the exact order used in the guide (ATT01 x9, duplicate, ATT02 x4, ATT03 x3, ATT99)
rows = []
for code, d, lab, sh, ev in DAYS:
    for off, tm in ev:
        if off == 0:
            rows.append([code, d, tm, "DEV-01"])
        else:  # next day (over-midnight exit)
            y, m, dd = d.split("/")
            rows.append([code, "%s/%s/%02d" % (y, m, int(dd) + 1), tm, "DEV-01"])
# duplicate of the very first ATT01 event goes after the ATT01 block
att01 = [r for r in rows if r[0] == "ATT01"]
rest = [r for r in rows if r[0] != "ATT01"]
rows = att01 + [list(att01[0])] + rest + [["ATT99", "1405/07/04", "08:00", "DEV-01"]]
assert len(rows) == 18 and rows[9] == rows[0] and rows[17][0] == "ATT99"
mk("clock-import-sample.xlsx", ["کد پرسنلی", "تاریخ", "ساعت", "دستگاه"], rows, [14, 14, 10, 12], "ترددها")

legacy = [
 ["12/ص/1/00045", "1403/05/17", "صادره", "درخواست پیشنهاد قیمت قطعات آزمایشی", "شرکت مشتری آزمایشی ۱", "scan1.pdf|scan2.pdf"],
 ["131", "1403/08/02", "وارده", "پاسخ استعلام آزمایشی", "سازمان آزمایشی", "scan3.pdf"],
 ["03/د/1/00007", "1402/11/20", "داخلی", "اعلام برنامه تعمیرات آزمایشی", "مدیریت تولید", ""],
 ["12/ص/1/00046", "1403/05/30", "صادره", "ارسال مدارک آزمایشی", "شرکت مشتری آزمایشی ۱", "scan4.pdf"],
 ["132", "1403/09/10", "وارده", "دعوت‌نامه آزمایشی", "سازمان آزمایشی", ""],
]
with open(D + "legacy-letters-sample.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["شماره", "تاریخ", "نوع", "موضوع", "طرف مکاتبه", "فایل‌ها"])
    w.writerows(legacy)
print("data ok", len(rows))
