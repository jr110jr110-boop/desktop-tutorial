# -*- coding: utf-8 -*-
import os, sys, re, openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import _KEYS, T, Q_DASH, Q_FORGE, Q_CASH, Q_FORCE, Q_REPEAT, Q_HI, Q_STOCK, ITEM1, ITEM2
from openpyxl.styles import Font, Alignment, PatternFill
D = "/home/claude/kita1/out/PMIS_Local_AI_Test/3-Data/"
os.makedirs(D, exist_ok=True)
SRC = "/home/claude/backend/src/Services/NasbNiro/Ai/"

def mk(name, sheets, lines):
    wb = openpyxl.Workbook(); first = True
    for title, header, rows, widths in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title; ws.sheet_view.rightToLeft = True
        ws.append(header)
        for r in rows: ws.append(r)
        for c in ws[1]:
            c.font = Font(name="B Zar", bold=True, size=12); c.fill = PatternFill("solid", fgColor="DDDDDD"); c.alignment = Alignment(horizontal="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="B Zar", size=12); c.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(c.value, str): c.number_format = "@"
        for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"
    note = wb.create_sheet("توضیح"); note.sheet_view.rightToLeft = True
    note.column_dimensions["A"].width = 110
    for line in lines: note.append([line])
    for row in note.iter_rows():
        for c in row: c.font = Font(name="B Zar", size=12); c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(D + name)

# ---- 1. evaluation questions (parsed from the program code)
tools = dict(re.findall(r'Name = "(\w+)", Title = "([^"]+)"', open(SRC + "AiTools.cs", encoding="utf-8").read()))
ev = open(SRC + "AiEval.cs", encoding="utf-8").read()
cases = re.findall(r'^\s*C\("(\d+)",\s*"((?:[^"\\]|\\.)*)",\s*"([^"]*)"(?:,\s*"([^"]*)")?\)', ev, re.M)
assert len(cases) == 66, len(cases)
groups = [(1, 10, "مدیریت"), (11, 21, "مالی"), (22, 27, "انبار و موجودی"), (28, 31, "تدارکات"), (32, 36, "پروژه‌ها"),
          (37, 41, "کیفیت"), (42, 47, "فروش، خدمات و تولید"), (48, 54, "اسناد"), (55, 61, "بدون ابزار"), (62, 66, "ترکیبی")]
def grp(n):
    for a, b, g in groups:
        if a <= n <= b: return g
rows = []
for num, q, tl, facts in cases:
    names = [x.strip() for x in tl.split(",") if x.strip()]
    rows.append(["EV-" + num, grp(int(num)), q.replace('\\"', '"'), "، ".join(tools.get(x, x) for x in names) or "—", "ندارد" if not facts else "دارد"])
mk("eval-questions-list.xlsx", [("پرسش‌های ارزیابی", ["شماره", "گروه", "پرسش", "ابزارهای مورد انتظار", "عدد مرجع"], rows, [10, 22, 60, 40, 12])],
   ["فقط مرجع است؛ این فایل وارد برنامه نمی‌شود و ارزیابی همین ۶۶ مورد را از داخل خود برنامه اجرا می‌کند (مدیریت هوش مصنوعی ← ارزیابی).",
    "فهرست از کد برنامه خوانده شده است. «عدد مرجع: دارد» یعنی پاسخ درست باید عددی از خروجی ابزار را بدون تغییر بگوید.",
    "معیار پذیرش طراحی برای مدل واقعی: دست‌کم ۹۰٪ قبولی و صفر شکست عددی. شمار قبولی روی مدل آزمایشی فقط بررسی فنی است."])

# ---- 2. test data
items = [[ITEM1, "قطعه آزمون هوش مصنوعی ۱", "عدد", 50, 20, "1,000,000", 30], [ITEM2, "قطعه آزمون هوش مصنوعی ۲", "عدد", 10, 100, "500,000", 0]]
fixed = [[Q_DASH, "ابزار داشبورد مدیرعامل", "طبق نتایج سامانه: سه عدد اول خروجی ابزار", "مدیرعامل یا PMO"],
         [Q_FORGE, "پاسخ ساختگی با مبلغ 987654321", "رد می‌شود؛ قالب سامانه «پاسخ مدل قابل راستی‌آزمایی نبود؛ نتایج مستقیم سامانه:»", "همه"],
         [Q_CASH, "ابزار وضعیت نقدینگی", "برای نقش‌های مالی: طبق نتایج سامانه؛ برای انباردار: «این ابزار برای شما در دسترس نیست.»", "همه"],
         [Q_FORCE, "مدل ابزار را بی‌اجازه صدا می‌زند", "«شما به این اطلاعات دسترسی ندارید.»", "همه"],
         [Q_REPEAT, "بدون ابزار؛ عددهای نوبت قبل", "همان عددها", "همه"],
         [Q_HI, "پاسخ عمومی", "سلام؛ چه کمکی از من برمی‌آید؟", "همه"],
         [Q_STOCK, "ابزار اقلام زیر حداقل موجودی (کلیدواژه «موجودی»)", "طبق نتایج سامانه: عددها", "همه"]]
kw = [["خلاصه داشبورد مدیرعامل", "ceo_dashboard"], ["جعل", "پاسخ ساختگی با مبلغ 987654321"], ["نقدینگی", "cash_position"],
      ["اجبار", "cash_position حتی اگر ارائه نشده باشد"], ["موجودی", "items_below_minimum"], ["سند / روش", "search_documents"],
      ["تکرار", "تکرار عددهای قبلی بدون ابزار"], ["یادداشت مقایسه", "مقایسه‌ی پیشنهادها (دستیار بازرگانی)"], ["پیشنهاد دسته‌بندی", "دسته‌بندی پیشنهادی: Major؛ روش ریشه‌یابی: پنج چرا"],
      ["پیش‌نویس پاسخ", "پیش‌نویس پاسخ نامه"], ["خلاصه صورت‌جلسه", "خلاصه صورت‌جلسه"], ["گزارش روزانه", "ceo_dashboard"], ["روایت هفتگی", "project_portfolio"], ["غیر از اینها", "سلام؛ چه کمکی از من برمی‌آید؟"]]
seats = [["CEO-BRIEF", "گزارش روزانه‌ی مدیرعامل", "مدیرعامل", "هر روز ساعت 07:00", "بله"], ["PMO-WEEKLY", "روایت گزارش هفتگی پروژه‌ها", "معاونت PMO", "هر هفته پنجشنبه ساعت 07:00", "بله"],
         ["WH-REORDER", "پیش‌نویس درخواست خرید اقلام زیر حداقل", "انبار", "هر روز ساعت 06:30", "خیر"], ["PROC-COMPARE", "یادداشت مقایسه‌ی استعلام‌ها", "بازرگانی", "هر 15 دقیقه", "بله"],
         ["QC-NCR-ASSIST", "دستیار گزارش عدم انطباق", "کنترل کیفیت", "هر 10 دقیقه", "بله"], ["FIN-ANOMALY", "بازبینی ناهنجاری‌های مالی", "مالی", "هر روز ساعت 06:00", "خیر"],
         ["OFFICE-DRAFT", "پیش‌نویس پاسخ نامه و خلاصه‌ی صورت‌جلسه", "دبیرخانه", "هر 15 دقیقه", "بله"]]
vouch = [["سند رُند", "امروز", "10301 (تفصیلی مشتری آزمون) بدهکار / 60102 بستانکار", "2,000,000,000", "یادداشت «مبلغ رُند و بزرگ»"],
         ["سند غیررند", "امروز", "همان حساب‌ها", "2,050,000,000", "یادداشت ندارد"],
         ["سند تکراری ۱", "امروز", "همان حساب‌ها", "123,456,000", "قطعی‌شده با کنترل نفر دوم"],
         ["سند تکراری ۲", "امروز منهای ۲ روز", "همان حساب‌ها", "123,456,000", "یادداشت «سند یا ردیف تکراری»"],
         ["سند جمعه", "آخرین جمعه گذشته", "همان حساب‌ها", "77,000,000", "یادداشت «ثبت در روز تعطیل»"]]
mk("test-data-and-mock-questions.xlsx", [
    ("کالاها", ["کد", "عنوان", "واحد", "نقطه سفارش", "موجودی اول دوره", "فی", "مقدار پیش‌نویس انبار"], items, [20, 30, 10, 14, 16, 14, 18]),
    ("اسناد مالی", ["عنوان", "تاریخ", "حساب‌ها", "مبلغ (ریال)", "نتیجه مورد انتظار"], vouch, [18, 20, 50, 18, 36]),
    ("پرسش‌های ثابت", ["پرسش دقیق", "کار مدل آزمایشی", "پاسخ مورد انتظار", "کاربر"], fixed, [40, 36, 60, 18]),
    ("کلیدواژه‌های مدل آزمایشی", ["واژه در پرسش", "رفتار"], kw, [28, 60]),
    ("هفت دستیار", ["کد", "عنوان", "نقش", "زمان‌بندی", "نیاز به مدل"], seats, [18, 40, 16, 28, 14]),
], ["فقط مرجع آزمون است؛ این فایل خودکار وارد برنامه نمی‌شود. مقدارها را دستی وارد کنید.",
    "مدل آزمایشی مدل واقعی نیست؛ بر اساس واژه‌های کلیدی پاسخ ثابت می‌دهد. گفتار به متن آن همیشه «خلاصه داشبورد مدیرعامل» می‌نویسد.",
    "همه کدها آزمایشی‌اند و نام شخص یا رمز در آن‌ها نیست."])
print("data ok", len(rows))
