# -*- coding: utf-8 -*-
import os, csv, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from example import DAYS
D = "/home/claude/kitoffice2/out/PMIS_Office_R2_Test/3-Data/"
os.makedirs(D, exist_ok=True)

def mk(name, header, rows, widths, sheet="Sheet1"):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = sheet; ws.sheet_view.rightToLeft = True
    ws.append(header)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font = Font(name="B Zar", bold=True, size=12); c.fill = PatternFill("solid", fgColor="DDDDDD"); c.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="B Zar", size=12)
            c.number_format = "@" if isinstance(c.value, str) else "General"
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(D + name)

# clock import: rows in the exact order used in the guide (ATT01 x9, duplicate, ATT02 x4, ATT03 x3, ATT99)
rows = []
for code, d, lab, sh, ev in DAYS:
    for off, tm in ev:
        if off == 0:
            rows.append([code, d, tm, "DEV-01"])
        else:  # next day (over-midnight exit)
            y, m, dd = d.split("/")
            rows.append([code, "%s/%s/%02d" % (y, m, int(dd) + 1), tm, "DEV-01"])
# duplicate of the very first ATT01 event goes after the ATT01 block
att01 = [r for r in rows if r[0] == "ATT01"]
rest = [r for r in rows if r[0] != "ATT01"]
rows = att01 + [list(att01[0])] + rest + [["ATT99", "1405/07/04", "08:00", "DEV-01"]]
assert len(rows) == 18 and rows[9] == rows[0] and rows[17][0] == "ATT99"
mk("clock-import-sample.xlsx", ["کد پرسنلی", "تاریخ", "ساعت", "دستگاه"], rows, [14, 14, 10, 12], "ترددها")

legacy = [
 ["12/ص/1/00045", "1403/05/17", "صادره", "درخواست پیشنهاد قیمت قطعات آزمایشی", "شرکت مشتری آزمایشی ۱", "scan1.pdf|scan2.pdf"],
 ["131", "1403/08/02", "وارده", "پاسخ استعلام آزمایشی", "سازمان آزمایشی", "scan3.pdf"],
 ["03/د/1/00007", "1402/11/20", "داخلی", "اعلام برنامه تعمیرات آزمایشی", "مدیریت تولید", ""],
 ["12/ص/1/00046", "1403/05/30", "صادره", "ارسال مدارک آزمایشی", "شرکت مشتری آزمایشی ۱", "scan4.pdf"],
 ["132", "1403/09/10", "وارده", "دعوت‌نامه آزمایشی", "سازمان آزمایشی", ""],
]
with open(D + "legacy-letters-sample.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["شماره", "تاریخ", "نوع", "موضوع", "طرف مکاتبه", "فایل‌ها"])
    w.writerows(legacy)

# ---------------------------------------------------------------- PW Kara export (R2)
# ATT04: 5 days x in/out (Mehr 1405: 07/12..07/15, 07/18) + the unknown code ATT99 (2 rows) = 12 rows. Persian dates as text yyyy/mm/dd.
PW = [  # (persian date, gregorian dd/mm/yyyy, [(time, direction)])
 ("1405/07/12", "04/10/2026", [("08:16", "ورود"), ("15:40", "خروج")]),
 ("1405/07/13", "05/10/2026", [("07:50", "ورود"), ("16:20", "خروج")]),
 ("1405/07/14", "06/10/2026", [("08:14", "ورود"), ("16:00", "خروج")]),
 ("1405/07/15", "07/10/2026", [("08:00", "ورود"), ("17:30", "خروج")]),
 ("1405/07/18", "10/10/2026", [("08:00", "ورود"), ("16:00", "خروج")]),
]
def pw_rows(greg):
    r = []
    for fa_d, gr_d, ev in PW:
        for tm, dr in ev:
            r.append(["ATT04", gr_d if greg else fa_d, tm, dr, "DEV-01"])
    d0 = PW[0][1] if greg else PW[0][0]
    r.append(["ATT99", d0, "08:00", "ورود", "DEV-01"])
    r.append(["ATT99", d0, "16:00", "خروج", "DEV-01"])
    return r
pw_head = ["کد پرسنلی", "تاریخ", "ساعت", "ورود/خروج", "دستگاه"]
mk("pw-kara-sample.xlsx", pw_head, pw_rows(False), [14, 14, 10, 12, 12], "PW Kara")
mk("pw-kara-ambiguous-dates.xlsx", pw_head, pw_rows(True), [14, 14, 10, 12, 12], "PW Kara")
assert len(pw_rows(False)) == 12

# ---------------------------------------------------------------- archive PDFs (R2): tiny valid single-page PDFs
def tiny_pdf(text):
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = "BT /F1 14 Tf 20 80 Td (%s) Tj ET" % text
    objs[3] = "<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += ("%d 0 obj\n%s\nendobj\n" % (i, o)).encode("latin-1")
    x = len(out)
    out += ("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode()
    for o in offs: out += ("%010d 00000 n \n" % o).encode()
    out += ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x)).encode()
    return out
AP = D + "archive-pdf/"
os.makedirs(AP, exist_ok=True)
for name, label in [("12_ص_1_00045.pdf", "Test scan 12-1-00045"), ("131.pdf", "Test scan 131"), ("03_د_1_00007.pdf", "Test scan 03-1-00007"), ("999.pdf", "Test scan 999 (no letter)")]:
    open(AP + name, "wb").write(tiny_pdf(label))
open(AP + "bad-scan.pdf", "wb").write("این فایل PDF نیست؛ فقط متن است.\n".encode("utf-8"))
print("data ok", len(rows))
