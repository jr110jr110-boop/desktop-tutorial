# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from build import T, SEC, QUESTIONS, LINKED, plain
OUT="/home/claude/kitoffice2/out/PMIS_Office_R2_Test/4-Feedback/Office-R2-Test-Feedback.xlsx"
os.makedirs(os.path.dirname(OUT),exist_ok=True)
F="B Zar"
thin=Side(style="thin",color="999999"); BR=Border(left=thin,right=thin,top=thin,bottom=thin)
HF=PatternFill("solid",fgColor="DDDDDD")
wb=openpyxl.Workbook()

def head(ws,cols,widths):
    ws.append(cols)
    for c in ws[1]:
        c.font=Font(name=F,bold=True,size=12); c.fill=HF; c.border=BR
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for i,w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
    ws.freeze_panes="A2"; ws.sheet_view.rightToLeft=True
    ws.row_dimensions[1].height=34

def body(ws,first_row,ncols,wrap_cols=None):
    for row in ws.iter_rows(min_row=first_row,max_row=ws.max_row,max_col=ncols):
        for c in row:
            c.font=Font(name=F,size=11); c.border=BR
            c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="right")

# ---- sheet 1
ws=wb.active; ws.title="نتایج تست"
cols=["شناسه","بخش","عنوان آزمون","نقش","نتیجه (OK / مشکل / انجام نشد)","توضیح (چه دیدید)","نام فایل تصویر"]
head(ws,cols,[9,30,40,26,22,60,24])
for x in T:
    ws.append([x["id"],SEC[x["sec"]],plain(x["title"]),plain(x["role"]),None,None,None])
n=ws.max_row
body(ws,2,7)
for r in range(2,n+1):
    ws.cell(r,1).alignment=Alignment(horizontal="center",vertical="top")
    ws.cell(r,1).font=Font(name=F,size=11,bold=True)
    ws.cell(r,5).alignment=Alignment(horizontal="center",vertical="top",wrap_text=True)
    for c in (5,6,7): ws.cell(r,c).fill=PatternFill("solid",fgColor="FFFFF2")
dv=DataValidation(type="list",formula1='"OK,مشکل,انجام نشد"',allow_blank=True,showErrorMessage=True,errorTitle="مقدار نامعتبر",error="فقط OK، مشکل یا انجام نشد")
dv.promptTitle="نتیجه"; dv.prompt="OK = مطابق نتیجه مورد انتظار؛ مشکل = نادرست (توضیح بنویسید)؛ انجام نشد = آزمون اجرا نشد (علت را بنویسید)"; dv.showInputMessage=True
ws.add_data_validation(dv); dv.add("E2:E%d"%n)
green=PatternFill("solid",bgColor="C6EFCE",fgColor="C6EFCE"); red=PatternFill("solid",bgColor="FFC7CE",fgColor="FFC7CE"); grey=PatternFill("solid",bgColor="E0E0E0",fgColor="E0E0E0")
ws.conditional_formatting.add("E2:E%d"%n,CellIsRule(operator="equal",formula=['"OK"'],fill=green))
ws.conditional_formatting.add("E2:E%d"%n,CellIsRule(operator="equal",formula=['"مشکل"'],fill=red))
ws.conditional_formatting.add("E2:E%d"%n,CellIsRule(operator="equal",formula=['"انجام نشد"'],fill=grey))
ws.auto_filter.ref="A1:G%d"%n
ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
ws.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
ws.print_title_rows="1:1"

# ---- sheet 2
w2=wb.create_sheet("پرسش‌های شرکت")
head(w2,["شماره","پرسش","آزمون مرتبط","پاسخ","تصمیم‌گیرنده (عنوان شغلی)","تاریخ (شمسی)"],[8,70,16,60,24,14])
for i,q in enumerate(QUESTIONS,1):
    w2.append([i,plain(q),plain(LINKED[i]).replace("، ","\n"),None,None,None])
body(w2,2,6)
for r in range(2,w2.max_row+1):
    w2.cell(r,1).alignment=Alignment(horizontal="center",vertical="top")
    w2.row_dimensions[r].height=95
    for c in (4,5,6): w2.cell(r,c).fill=PatternFill("solid",fgColor="FFFFF2")
w2.page_setup.orientation="landscape"; w2.page_setup.fitToWidth=1; w2.page_setup.fitToHeight=0
w2.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)

# ---- sheet 3
w3=wb.create_sheet("نظرهای کلی")
head(w3,["موضوع (منو یا بخش)","نظر یا پیشنهاد","اهمیت (زیاد / متوسط / کم)","نام فایل تصویر"],[30,90,24,24])
for _ in range(30): w3.append([None,None,None,None])
body(w3,2,4)
dv3=DataValidation(type="list",formula1='"زیاد,متوسط,کم"',allow_blank=True)
w3.add_data_validation(dv3); dv3.add("C2:C31")
for r in range(2,w3.max_row+1):
    w3.row_dimensions[r].height=40
    for c in (1,2,3,4): w3.cell(r,c).fill=PatternFill("solid",fgColor="FFFFF2")

# ---- sheet 4
w4=wb.create_sheet("راهنما")
w4.sheet_view.rightToLeft=True; w4.column_dimensions["A"].width=26; w4.column_dimensions["B"].width=100
rows=[
 ("هدف","ثبت نتیجه آزمون بسته اتوماسیون اداری PMIS (نسخه R2، پاسخ‌های شرکت). هر ردیف برگه «نتایج تست» یک آزمون از راهنمای PDF (0-Test-Guide-Office-R2-FA.pdf) است و کد آن (O-01 تا O-127) همان کد راهنماست. آزمون‌های بخش «R2 — پاسخ‌های شرکت» (O-107 تا O-127) تازه‌اند."),
 ("پایگاه آزمون","همه کارها فقط روی پایگاه آزمایشی انجام می‌شود، با داده و نام آزمایشی. رمز یا نام کاربری در این فرم نوشته نمی‌شود."),
 ("ستون نتیجه","فقط یکی از OK، مشکل یا انجام نشد را از فهرست انتخاب کنید. OK یعنی هر آنچه در «نتیجه مورد انتظار» راهنما آمده دیده شد (از جمله پیام‌ها و عددها). مشکل یعنی نتیجه فرق داشت یا خطا آمد. انجام نشد یعنی آزمون اجرا نشد."),
 ("توضیح","برای هر «مشکل» بنویسید چه دیدید (متن پیام خطا، عدد دیده‌شده در برابر عدد مورد انتظار، منوی محل). برای «انجام نشد» علت را بنویسید."),
 ("نام فایل تصویر","نام فایل تصویر (مثلاً O-26-1) که در پوشه 4-Feedback کنار همین فایل گذاشته‌اید."),
 ("تفکیک وظایف","در آزمون‌هایی که دو نفر لازم است با دو کاربر جدا کار کنید. اگر یک نفر توانست کار تنظیم و تایید را هر دو بکند، «مشکل» بزنید و بنویسید «بحرانی»."),
 ("نقش‌ها","عنوان شغلی، نه نام شخص: دبیرخانه، اداری و منابع انسانی، مدیرعامل، معاونت PMO، حقوق و دستمزد، کارشناس واحد، مدیر واحد."),
 ("تاریخ","تاریخ شمسی به‌صورت 1405/08/10. تاریخ میلادی نوشته نشود."),
 ("برگه پرسش‌های شرکت","پاسخ هر پرسش، عنوان شغلی تصمیم‌گیرنده و تاریخ را بنویسید."),
 ("پس از پایان","فایل را با نام Office-R2-Test-Feedback-(عنوان شغلی).xlsx ذخیره و برای مدیر فنی بفرستید."),
]
w4.append(["موضوع","توضیح"])
for c in w4[1]: c.font=Font(name=F,bold=True,size=12); c.fill=HF; c.border=BR; c.alignment=Alignment(horizontal="center")
for a,b in rows: w4.append([a,b])
for row in w4.iter_rows(min_row=2):
    for c in row:
        c.font=Font(name=F,size=11,bold=(c.column==1)); c.border=BR; c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="right")
wb.save(OUT); print("saved",n-1,"tests")
